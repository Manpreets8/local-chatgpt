import io

import pytest
from docx import Document as DocxDocument

from app.services import document_service


def test_clean_text_collapses_whitespace_and_blank_lines():
    dirty = "Hello   world\t\there\n\n\n\nNext paragraph"
    cleaned = document_service.clean_text(dirty)
    assert cleaned == "Hello world here\n\nNext paragraph"


def test_clean_text_strips_null_bytes():
    assert document_service.clean_text("a\x00b") == "ab"


def test_chunk_text_empty_returns_empty_list():
    assert document_service.chunk_text("") == []
    assert document_service.chunk_text("   ") == []


def test_chunk_text_short_text_returns_single_chunk():
    chunks = document_service.chunk_text("short text", chunk_size=800, overlap=100)
    assert chunks == ["short text"]


def test_chunk_text_splits_long_text_with_overlap():
    text = "a" * 1000
    chunks = document_service.chunk_text(text, chunk_size=400, overlap=50)
    assert len(chunks) > 1
    assert all(len(c) <= 400 for c in chunks)
    # overlap means the tail of one chunk reappears at the start of the next
    assert chunks[0][-50:] == chunks[1][:50]


def test_extract_pages_txt():
    pages = document_service.extract_pages(b"hello from a text file", "txt")
    assert pages == [(None, "hello from a text file")]


def test_extract_pages_docx():
    doc = DocxDocument()
    doc.add_paragraph("First paragraph.")
    doc.add_paragraph("Second paragraph.")
    buffer = io.BytesIO()
    doc.save(buffer)

    pages = document_service.extract_pages(buffer.getvalue(), "docx")
    assert len(pages) == 1
    page_number, text = pages[0]
    assert page_number is None
    assert "First paragraph." in text
    assert "Second paragraph." in text


def test_validate_upload_rejects_oversized_file():
    with pytest.raises(document_service.FileTooLargeError):
        document_service.validate_upload("text/plain", document_service.MAX_FILE_SIZE_BYTES + 1)


def test_validate_upload_rejects_unsupported_type():
    with pytest.raises(document_service.UnsupportedFileTypeError):
        document_service.validate_upload("application/zip", 100)


@pytest.mark.parametrize(
    "content_type,expected_kind",
    [
        ("application/pdf", "pdf"),
        ("text/plain", "txt"),
        (
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            "docx",
        ),
    ],
)
def test_validate_upload_accepts_supported_types(content_type, expected_kind):
    assert document_service.validate_upload(content_type, 100) == expected_kind


def test_build_chunks_preserves_page_numbers():
    pages = [(1, "Page one text."), (2, "Page two text.")]
    chunks = document_service.build_chunks(pages)
    assert chunks == [(1, "Page one text."), (2, "Page two text.")]
